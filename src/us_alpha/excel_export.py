# -*- coding: utf-8 -*-
"""
US Alpha Selection System - Excel Export.

Exports all database tables to a formatted Excel workbook on Desktop.
Each table becomes a separate sheet with:
- Column headers with filters
- Conditional formatting for grades
- Proper column widths
- Summary sheet with statistics

Output: ~/Desktop/us_alpha_db/US_Alpha_Records.xlsx
"""

import os
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .storage import AlphaDB


# ─── Style Constants ──────────────────────────────────────────────────────────

_HEADER_FONT = Font(name="Microsoft JhengHei", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_DATA_FONT = Font(name="Microsoft JhengHei", size=10)
_DATA_ALIGNMENT = Alignment(vertical="center", wrap_text=False)

_GRADE_FILLS = {
    "A": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "B": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "C": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "D": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "F": PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
    "P": PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid"),
}

_PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_WARN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_BLOCK_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


# ─── Sheet Definitions ────────────────────────────────────────────────────────

_SHEETS = {
    "推薦紀錄": {
        "table": "formal_rec",
        "columns": [
            ("rec_id", "推薦ID", 18),
            ("rec_date", "日期", 12),
            ("ticker", "股票", 8),
            ("company_name", "公司", 16),
            ("strategy_version", "策略版本", 20),
            ("momentum_state", "動能狀態", 10),
            ("rec_type", "推薦類型", 10),
            ("buy_price", "買入價", 10),
            ("stop_loss", "停損", 10),
            ("target_1", "目標1", 10),
            ("target_2", "目標2", 10),
            ("risk_reward_ratio", "RR比", 8),
            ("final_grade", "等級", 6),
            ("total_score", "總分", 6),
            ("market_state", "市場狀態", 14),
            ("position_risk", "倉位風險", 10),
            ("futu_priority", "富途優先", 10),
            ("mfe", "MFE%", 8),
            ("mae", "MAE%", 8),
            ("final_outcome", "最終結果", 12),
        ],
    },
    "評分審計": {
        "table": "score_audit",
        "columns": [
            ("score_id", "評分ID", 22),
            ("ticker", "股票", 8),
            ("timestamp", "時間", 18),
            ("strategy_version", "策略版本", 20),
            ("original_score", "原始分", 8),
            ("final_score", "最終分", 8),
            ("final_grade", "等級", 6),
            ("risk_lock", "風險封鎖", 10),
            ("enters_rec", "進入推薦", 10),
        ],
    },
    "觀察名單": {
        "table": "watchlist",
        "columns": [
            ("watch_id", "觀察ID", 18),
            ("date", "日期", 12),
            ("ticker", "股票", 8),
            ("layer", "分層", 10),
            ("watch_reason", "觀察原因", 30),
            ("fail_reason", "未通過原因", 30),
            ("trigger_condition", "觸發條件", 20),
            ("can_convert_to_rec", "可轉推薦", 10),
            ("next_check_date", "下次檢查", 12),
        ],
    },
    "市場環境": {
        "table": "market_gate",
        "columns": [
            ("date", "日期", 12),
            ("market_state", "市場狀態", 14),
            ("spy_trend", "SPY", 8),
            ("qqq_trend", "QQQ", 8),
            ("iwm_trend", "IWM", 8),
            ("smh_trend", "SMH", 8),
            ("vix_level", "VIX", 8),
            ("us10y_yield", "10Y殖利率", 10),
            ("notes", "備註", 30),
        ],
    },
    "板塊強度": {
        "table": "sector_rank",
        "columns": [
            ("date", "日期", 12),
            ("sector_name", "板塊", 28),
            ("total_score", "總分", 8),
            ("grade", "等級", 6),
            ("representative_tickers", "代表股", 20),
        ],
    },
    "催化劑": {
        "table": "catalyst_calendar",
        "columns": [
            ("catalyst_id", "催化劑ID", 22),
            ("ticker", "股票", 8),
            ("event_type", "事件類型", 12),
            ("expected_date", "預計日期", 12),
            ("certainty", "確定性", 8),
        ],
    },
    "資訊收錄": {
        "table": "info_inbox",
        "columns": [
            ("info_id", "資訊ID", 18),
            ("timestamp", "時間", 18),
            ("market", "市場", 6),
            ("tickers", "涉及股票", 14),
            ("source_grade", "來源等級", 10),
            ("initial_direction", "方向", 10),
            ("enters_scoring", "進入主分數", 10),
            ("enters_market_gate", "進入市場閘門", 12),
            ("enters_risk_lock", "進入風險封鎖", 12),
            ("enters_sector_rank", "進入板塊分數", 12),
            ("conclusion", "結論", 8),
            ("raw_content", "內容摘要", 50),
        ],
    },
    "排除名單": {
        "table": "exclusion",
        "columns": [
            ("ticker", "股票", 8),
            ("date", "日期", 12),
            ("reason", "排除原因", 40),
            ("category", "類別", 12),
            ("duration", "期限", 12),
        ],
    },
    "持倉管理": {
        "table": "holding",
        "columns": [
            ("ticker", "股票", 8),
            ("timestamp", "時間", 18),
            ("action", "動作", 10),
            ("shares", "股數", 8),
            ("avg_cost", "成本", 10),
            ("current_price", "現價", 10),
            ("position_value", "持倉金額", 12),
            ("pnl_pct", "盈虧%", 8),
            ("linked_rec_id", "推薦ID", 16),
            ("original_buy_reason", "買入理由", 20),
            ("alert_price", "警戒價", 8),
            ("invalidation_price", "失效價", 8),
            ("stop_loss", "停損價", 8),
            ("target_1", "T1", 8),
            ("target_2", "T2", 8),
            ("can_add", "可否加碼", 10),
            ("add_condition", "加碼條件", 16),
            ("holding_days", "持有天數", 10),
            ("original_thesis_valid", "邏輯有效", 10),
            ("next_action", "下一步", 12),
        ],
    },
    "復盤紀錄": {
        "table": "postmortem",
        "columns": [
            ("rec_id", "推薦ID", 18),
            ("ticker", "股票", 8),
            ("review_date", "復盤日期", 12),
            ("review_period", "週期", 8),
            ("actual_performance_pct", "實際報酬%", 10),
            ("mfe", "MFE%", 8),
            ("mae", "MAE%", 8),
            ("error_type", "錯誤類型", 16),
            ("needs_model_change", "需改模型", 10),
            ("conclusion", "結論", 30),
        ],
    },
    "交易計畫": {
        "table": "trade_plan",
        "columns": [
            ("rec_id", "推薦ID", 18),
            ("ticker", "股票", 8),
            ("date", "日期", 12),
            ("strategy_version", "策略", 16),
            ("market_state", "市場", 12),
            ("entry_trigger", "進場觸發", 20),
            ("stop_loss", "停損", 10),
            ("target_1", "目標1", 10),
            ("target_2", "目標2", 10),
            ("position_size", "倉位", 10),
            ("timing_gate", "時間閘門", 14),
            ("exit_rules", "出場規則", 20),
            ("counter_thesis", "反方論述", 30),
            ("plan_text", "完整計畫", 50),
        ],
    },
    "時間閘門": {
        "table": "timing_gate",
        "columns": [
            ("ticker", "股票", 8),
            ("date", "日期", 12),
            ("timing_state", "時間狀態", 12),
            ("earnings_date", "財報日期", 12),
            ("days_to_earnings", "距財報天數", 10),
            ("catalyst_date", "催化劑日期", 12),
            ("days_to_catalyst", "距事件天數", 10),
            ("is_fomc_week", "FOMC週", 8),
            ("is_opex_week", "OPEX週", 8),
            ("is_cpi_nfp_week", "CPI/NFP", 8),
            ("is_earnings_season", "財報季", 8),
            ("timing_passed", "通過", 6),
            ("timing_action", "動作", 14),
            ("notes", "備註", 30),
        ],
    },
    "出場計畫": {
        "table": "exit_plan",
        "columns": [
            ("ticker", "股票", 8),
            ("rec_id", "推薦ID", 16),
            ("date", "日期", 12),
            ("entry_price", "進場價", 10),
            ("stop_loss", "停損", 10),
            ("trailing_stop_rule", "移動停損規則", 20),
            ("target_1", "T1", 10),
            ("t1_exit_pct", "T1出場比", 10),
            ("target_2", "T2", 10),
            ("t2_exit_pct", "T2出場比", 10),
            ("remainder_rule", "剩餘部位", 16),
            ("time_stop_days", "時間停損天數", 12),
            ("catalyst_fail_rule", "利多出盡規則", 20),
            ("structure_break_rule", "結構失效規則", 20),
            ("profit_lock_rules", "獲利鎖定規則", 24),
            ("actual_exit_trigger", "實際出場觸發", 14),
            ("actual_exit_price", "實際出場價", 12),
            ("actual_exit_date", "實際出場日", 12),
        ],
    },
    "模型修正": {
        "table": "model_change",
        "columns": [
            ("change_date", "日期", 12),
            ("module_changed", "修改模組", 16),
            ("change_description", "修改內容", 40),
            ("rationale", "原因", 30),
            ("triggered_by", "觸發來源", 16),
        ],
    },
    "環境記憶": {
        "table": "regime_memory",
        "columns": [
            ("regime_id", "環境ID", 18),
            ("market_state", "市場狀態", 14),
            ("start_date", "開始日期", 12),
            ("end_date", "結束日期", 12),
            ("trigger_reason", "觸發原因", 30),
            ("spy_state", "SPY", 8),
            ("qqq_state", "QQQ", 8),
            ("vix_change", "VIX變化", 10),
            ("lessons", "經驗教訓", 40),
        ],
    },
}


def _default_export_path() -> Path:
    env_path = os.environ.get("US_ALPHA_DB_PATH")
    if env_path:
        return Path(env_path).parent / "US_Alpha_Records.xlsx"
    return Path.home() / "Desktop" / "us_alpha_db" / "US_Alpha_Records.xlsx"


def _apply_header_style(ws, col_count: int):
    """Apply header formatting to first row."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _THIN_BORDER


def _apply_grade_formatting(ws, col_idx: int, max_row: int):
    """Apply conditional color to grade columns."""
    col_letter = get_column_letter(col_idx)
    for row in range(2, max_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        val = str(cell.value or "").strip()
        if val in _GRADE_FILLS:
            cell.fill = _GRADE_FILLS[val]


def _apply_risk_formatting(ws, col_idx: int, max_row: int):
    """Apply conditional color to risk/lock columns."""
    col_letter = get_column_letter(col_idx)
    for row in range(2, max_row + 1):
        cell = ws.cell(row=row, column=col_idx)
        val = str(cell.value or "").lower().strip()
        if val == "pass":
            cell.fill = _PASS_FILL
        elif val == "warning":
            cell.fill = _WARN_FILL
        elif val == "blocked":
            cell.fill = _BLOCK_FILL


def _write_sheet(ws, rows: List[Dict], columns: list):
    """Write data rows to a worksheet with formatting."""
    for col_idx, (key, header, width) in enumerate(columns, 1):
        ws.cell(row=1, column=col_idx, value=header)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    _apply_header_style(ws, len(columns))

    grade_cols = []
    risk_cols = []

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, (key, header, width) in enumerate(columns, 1):
            value = row_data.get(key)
            if value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _DATA_FONT
            cell.alignment = _DATA_ALIGNMENT
            cell.border = _THIN_BORDER

    max_row = len(rows) + 1

    for col_idx, (key, header, width) in enumerate(columns, 1):
        if "grade" in key or key == "grade":
            _apply_grade_formatting(ws, col_idx, max_row)
        elif "risk_lock" in key or "risk" in key.lower():
            _apply_risk_formatting(ws, col_idx, max_row)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max_row}"
    ws.freeze_panes = "A2"


def _write_summary_sheet(ws, db: AlphaDB):
    """Write summary/dashboard sheet."""
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 30

    title_font = Font(name="Microsoft JhengHei", bold=True, size=14)
    section_font = Font(name="Microsoft JhengHei", bold=True, size=12)
    data_font = Font(name="Microsoft JhengHei", size=11)

    row = 1
    ws.cell(row=row, column=1, value="美股高報酬交易系統").font = title_font
    row += 1
    ws.cell(row=row, column=1, value=f"匯出時間：{datetime.now().strftime('%Y-%m-%d %H:%M')}").font = data_font
    row += 2

    ws.cell(row=row, column=1, value="資料表統計").font = section_font
    row += 1
    counts = db.get_table_counts()
    table_names_zh = {
        "info_inbox": "資訊收錄",
        "market_gate": "市場環境",
        "sector_rank": "板塊強度",
        "catalyst_calendar": "催化劑",
        "score_audit": "評分審計",
        "formal_rec": "正式推薦",
        "watchlist": "觀察名單",
        "exclusion": "排除名單",
        "holding": "持倉管理",
        "postmortem": "復盤紀錄",
        "model_change": "模型修正",
        "trade_plan": "交易計畫",
        "timing_gate": "時間閘門",
        "exit_plan": "出場計畫",
        "regime_memory": "環境記憶",
    }
    for table, count in counts.items():
        zh_name = table_names_zh.get(table, table)
        ws.cell(row=row, column=1, value=zh_name).font = data_font
        ws.cell(row=row, column=2, value=count).font = data_font
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="勝率統計").font = section_font
    row += 1
    stats = db.get_win_rate()
    stat_labels = [
        ("總推薦數", stats.get("total", 0)),
        ("勝利", stats.get("wins", 0)),
        ("部分成功", stats.get("partial", 0)),
        ("虧損", stats.get("losses", 0)),
        ("勝率", f"{stats.get('win_rate', 0)}%"),
        ("平均報酬", f"{stats.get('avg_return', 0)}%"),
        ("平均MFE", f"{stats.get('avg_mfe', 0)}%"),
        ("平均MAE", f"{stats.get('avg_mae', 0)}%"),
    ]
    for label, value in stat_labels:
        ws.cell(row=row, column=1, value=label).font = data_font
        ws.cell(row=row, column=2, value=value).font = data_font
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="系統版本").font = section_font
    row += 1
    ws.cell(row=row, column=1, value="HIGH_ALPHA_LAUNCHPAD_US_v2.2").font = data_font
    row += 1
    ws.cell(row=row, column=1, value="含：TIMING + EXIT + TRADE_PLAN").font = data_font


def export_to_excel(db: Optional[AlphaDB] = None, output_path: Optional[str] = None) -> str:
    """
    Export all database tables to a formatted Excel workbook.

    Args:
        db: AlphaDB instance (creates one if not provided)
        output_path: Custom output path (defaults to ~/Desktop/us_alpha_db/)

    Returns:
        Path to the created Excel file.
    """
    close_db = False
    if db is None:
        db = AlphaDB()
        close_db = True

    path = Path(output_path) if output_path else _default_export_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "總覽"
    _write_summary_sheet(summary_ws, db)

    for sheet_name, config in _SHEETS.items():
        table = config["table"]
        columns = config["columns"]

        rows = db._conn.execute(f"SELECT * FROM {table} ORDER BY id DESC").fetchall()
        rows_data = [dict(r) for r in rows]

        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, rows_data, columns)

    wb.save(str(path))

    if close_db:
        db.close()

    return str(path)
